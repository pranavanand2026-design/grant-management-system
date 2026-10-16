"""
Parser for client cash flow statements exported from accounting software.

These have no fixed template. The only structure we rely on: some sheet has
a header row containing several date-typed cells (one per month), and below
it rows with a text label in an early column and numbers under the month
columns. Each (account, month) pair becomes one row for the tagging table —
the unit a grant manager tags against a budget category.

Subtotal-ish rows (total / net / opening / closing balances) are skipped so
tagging never double-counts. Everything else is kept, untagged; untagged is
not an error state. Every row is imported — nothing is ever held out, only
flagged for her to decide on (see suspected_rollup below).
"""

from __future__ import annotations

import datetime as _dt
import re

import openpyxl

_SKIP = re.compile(r"^(total|net |opening|closing|cash flow|balance)", re.IGNORECASE)

# Loose on purpose. This used to be a narrow pattern aimed at catching only an
# accountant's own "Eligible Expenditure: <category>" rollup line, and it was
# wrong in both directions on the sample file it was tuned against: it held
# out "Ineligible expenditure: Purchase machinery" — a real
# $690k capital purchase, the largest line in the file — and it missed
# "Eligible Exenditure: Contract" because the source's own typo doesn't match
# any tidy alternation. A regex cannot reliably tell "the preparer's derived
# figure" from "a real account whose name mentions eligibility" — that's a
# judgment call, not a structural fact. So this only ever flags a *question*;
# nothing it matches is ever excluded from `rows`.
_SUSPECT = re.compile(r"eligible", re.IGNORECASE)


def _is_date(v) -> bool:
    return isinstance(v, (_dt.datetime, _dt.date))


def extract(path: str) -> dict:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise ValueError("This file could not be read as an Excel workbook.") from exc

    best = None  # (sheet, header_row_idx, [(col, iso_month)])
    for ws in wb.worksheets:
        for ri in range(1, min(8, ws.max_row) + 1):
            cols = [
                (ci, ws.cell(ri, ci).value)
                for ci in range(1, min(ws.max_column, 60) + 1)
                if _is_date(ws.cell(ri, ci).value)
            ]
            if len(cols) >= 3 and (best is None or len(cols) > len(best[2])):
                months = [(ci, f"{v.year:04d}-{v.month:02d}") for ci, v in cols]
                best = (ws, ri, months)
    if best is None:
        raise ValueError(
            "No month columns found. This does not look like a cash flow "
            "statement export (expected a header row of month dates)."
        )

    ws, header_row, months = best
    rows, row_id = [], 0
    current_section = None
    for ri in range(header_row + 1, ws.max_row + 1):
        label = None
        for ci in range(1, months[0][0]):
            v = ws.cell(ri, ci).value
            if isinstance(v, str) and v.strip():
                label = re.sub(r"\s+", " ", v).strip()
                break
        if not label or _SKIP.match(label):
            continue

        values = [ws.cell(ri, ci).value for ci, _period in months]
        has_amount = any(isinstance(v, (int, float)) and v != 0 for v in values)
        if not has_amount:
            # A labelled row with no figure under any period is a section
            # divider ("Cash incoming", "Cash outgoing"), not an account —
            # carry its text forward, verbatim, onto the rows beneath it.
            current_section = label
            continue

        suspected_rollup = bool(_SUSPECT.search(label))
        for ci, period in months:
            v = ws.cell(ri, ci).value
            if isinstance(v, (int, float)) and v != 0:
                row_id += 1
                rows.append({
                    "id": f"sr_{row_id}",
                    "period": period,
                    "account": label,
                    "amount": round(float(v), 2),
                    "section": current_section,
                    "suspectedRollup": suspected_rollup,
                })
    if not rows:
        raise ValueError("Month columns were found but no account rows beneath them.")

    periods = sorted({r["period"] for r in rows})
    return {
        "source": "statement",
        "sheet": ws.title,
        "periods": periods,
        "period_start": periods[0],
        "period_end": periods[-1],
        "rows": rows,
        "total": round(sum(r["amount"] for r in rows), 2),
    }
